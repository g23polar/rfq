import importlib.util
import json
from pathlib import Path
from typing import get_args

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from rfq.models import PRICED_STATUSES, IssueKind, LineStatus
from rfq.reference import load_reference

ROOT = Path(__file__).parent.parent
FIXTURES = ROOT / "fixtures"

# M cases of the spec's edge-case catalogue (section 12) that a fixture must carry.
# S2 (leak check) and S3 (question promises) are behaviors with their own unit tests in task 12.
M_CASES = {"D1", "D2", "D5", "D7", "D10", "L1", "L2", "L5", "U1", "U2", "U4", "U6", "U7", "M1",
           "M2", "M5", "Q1", "Q2", "Q4", "P1", "P2", "P3", "P5", "X1", "X3", "X4", "S1", "B1"}


@pytest.fixture(scope="module")
def gen():
    spec = importlib.util.spec_from_file_location("make_fixtures",
                                                  ROOT / "scripts" / "make_fixtures.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def golden() -> list[dict]:
    return [json.loads(l) for l in (FIXTURES / "golden.jsonl").read_text().splitlines()]


def pdf_text(name: str) -> str:
    return "\n".join(p.extract_text() for p in PdfReader(FIXTURES / name).pages)


def test_golden_matches_generator(gen, golden):
    assert golden == json.loads(json.dumps(gen.GOLDEN, ensure_ascii=False))
    assert (FIXTURES / "rfq_email.txt").read_text() == gen.email_text()


def test_golden_covers_every_status_and_m_case(gen, golden):
    assert len(golden) >= 25
    assert {r["expected"]["status"] for r in golden} == set(get_args(LineStatus))
    covered = {c for r in golden for c in r["case_ids"]}
    covered |= {c for cases in gen.FIXTURE_CASES.values() for c in cases}
    assert M_CASES <= covered, M_CASES - covered


def test_golden_records_are_consistent_with_the_catalog(golden):
    ref = load_reference()
    assert len({r["id"] for r in golden}) == len(golden)
    for r in golden:
        exp = r["expected"]
        assert all(ref.item(s) for s in exp["skus"]), r["id"]
        assert set(exp["issue_kinds"]) <= set(get_args(IssueKind)), r["id"]
        assert bool(exp["skus"]) == (exp["status"] in PRICED_STATUSES), r["id"]
        if exp["status"] == "quoted_with_alternates":
            assert not r["determinable"], r["id"]


def test_fixture_lines_are_verbatim_in_their_documents(golden):
    docs = {"rfq_email.txt": (FIXTURES / "rfq_email.txt").read_text(), "rfq.pdf": pdf_text("rfq.pdf")}
    for r in golden:
        if r["fixture"]:
            assert r["text"] in docs[r["fixture"]], r["id"]


def test_email_keeps_history_and_signature_below_the_request():
    text = (FIXTURES / "rfq_email.txt").read_text()
    request, history = text.split("-----Original Message-----")
    assert "Cell 212-555-0148" in request and "@ $68.40 ea" in history


def test_pdf_injection_is_the_only_difference(gen):
    injected, clean = pdf_text("rfq.pdf"), pdf_text("rfq_clean.pdf")
    assert gen.INJECTION in injected and gen.INJECTION not in clean
    assert injected.replace(gen.INJECTION, "").strip() == clean.strip()
    table_end = injected.index("Notes:")
    assert injected.index(gen.INJECTION) > table_end  # outside the line items


def test_spreadsheet_keeps_buyer_line_numbers():
    rows = list(load_workbook(FIXTURES / "rfq.xlsx").active.iter_rows(values_only=True))
    header = rows.index(("Line", "Description", "Qty", "UoM", "Notes"))
    assert [r[0] for r in rows[header + 1:]] == [10, 20, 30]
    assert rows[header + 2][2] is None  # Q1: missing quantity
